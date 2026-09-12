#!/usr/bin/env python3
"""Generic DOCX renderer (template `docx-generic`, python-docx) for advisory
deliverables and threat-intel briefs that no registered template governs.

Usage: render.py <data.json> <out.docx>

data.json contract (templates/docx_generic.schema.json):
{
  "title": str, "subtitle": str?, "supplier": str?, "service": str?,
  "reportDate": "YYYY-MM-DD"?, "classification": str?,   # default "Euronext Internal"
  "author": str?,                                          # producing agent
  "sections": [ { "heading": str, "level": 1|2|3?,
                  "paragraphs": [str, ...]?,               # **bold** inline supported
                  "bullets": [str, ...]?,
                  "table": {"columns": [str], "rows": [[...]]}? } ],
  "sources": [str, ...]?
}
House style (agents/document_agents_addendum.md): Verdana, teal
RGB(0,141,127) headings, classification in the header, title/date/sources.
"""

from __future__ import annotations

import json
import re
import sys
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

TEAL = RGBColor(0x00, 0x8D, 0x7F)
FONT = "Verdana"


def _runs(p, text: str, size=10):
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        r = p.add_run(part[2:-2] if bold else part)
        r.bold, r.font.name, r.font.size = bold, FONT, Pt(size)


def _heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb, r.font.name = TEAL, FONT
    return h


def _table(doc, spec: dict):
    cols = spec["columns"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for i, c in enumerate(cols):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(str(c))
        r.bold, r.font.name, r.font.size, r.font.color.rgb = True, FONT, Pt(9), TEAL
    for row in spec.get("rows", []):
        cells = t.add_row().cells
        for i, v in enumerate(row[:len(cols)]):
            cells[i].text = ""
            _runs(cells[i].paragraphs[0], "" if v is None else str(v), 9)


def render(d: dict, out: str) -> None:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(10)
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = f"{d.get('classification', 'Euronext Internal')} — Euronext InfoSec Assurance"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:
        r.font.size, r.font.name, r.font.color.rgb = Pt(8), FONT, RGBColor(0x66, 0x66, 0x66)
    t = doc.add_paragraph()
    _runs(t, f"**{d['title']}**", 18)
    for r in t.runs:
        r.font.color.rgb = TEAL
    if d.get("subtitle"):
        _runs(doc.add_paragraph(), d["subtitle"], 12)
    meta = [("Supplier", d.get("supplier")), ("Service", d.get("service")),
            ("Date", d.get("reportDate") or date.today().isoformat()),
            ("Produced by", d.get("author")), ("Classification", d.get("classification", "Euronext Internal"))]
    mt = doc.add_table(rows=0, cols=2)
    mt.style = "Table Grid"
    for k, v in meta:
        if v:
            c = mt.add_row().cells
            c[0].text, c[1].text = "", ""
            _runs(c[0].paragraphs[0], f"**{k}**", 9)
            _runs(c[1].paragraphs[0], str(v), 9)
    for s in d.get("sections", []):
        _heading(doc, s["heading"], int(s.get("level", 1)))
        for p in s.get("paragraphs", []):
            _runs(doc.add_paragraph(), p)
        for b in s.get("bullets", []):
            _runs(doc.add_paragraph(style="List Bullet"), b)
        if s.get("table"):
            _table(doc, s["table"])
    if d.get("sources"):
        _heading(doc, "Sources", 1)
        for i, src in enumerate(d["sources"], 1):
            _runs(doc.add_paragraph(style="List Number"), src, 9)
    doc.core_properties.title = d["title"]
    doc.core_properties.author = d.get("author", "Euronext InfoSec Assurance")
    doc.save(out)


if __name__ == "__main__":
    with open(sys.argv[1], encoding="utf-8") as fh:
        render(json.load(fh), sys.argv[2])
    print("OK", sys.argv[2])
