#!/usr/bin/env python3
"""Generic PPTX renderer (template `pptx-generic`, python-pptx) for advisory
decks that no registered template governs (16:9, Verdana, teal accents).

Usage: render.py <data.json> <out.pptx>

data.json contract (templates/pptx_generic.schema.json):
{ "title": str, "subtitle": str?, "reportDate": "YYYY-MM-DD"?, "classification": str?,
  "slides": [ { "title": str, "bullets": [str]?, "table": {"columns": [str], "rows": [[...]]}?,
                "notes": str? } ], "sources": [str]? }
"""

from __future__ import annotations

import json
import sys
from datetime import date

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

TEAL = RGBColor(0x00, 0x8D, 0x7F)
GREY = RGBColor(0x64, 0x74, 0x8B)
FONT = "Verdana"


def _text(tf, text, size=14, bold=False, color=None):
    p = tf.paragraphs[0] if not tf.paragraphs[0].text and len(tf.paragraphs) == 1 else tf.add_paragraph()
    p.text = text
    for r in p.runs:
        r.font.name, r.font.size, r.font.bold = FONT, Pt(size), bold
        if color:
            r.font.color.rgb = color
    return p


def _footer(slide, prs, d, n):
    tb = slide.shapes.add_textbox(Inches(0.4), prs.slide_height - Inches(0.45), prs.slide_width - Inches(0.8), Inches(0.3))
    _text(tb.text_frame, f"Euronext InfoSec Assurance · {d.get('classification', 'Euronext Internal')} · "
                         f"{d.get('reportDate') or date.today().isoformat()} · {n}", 9, color=GREY)


def render(d: dict, out: str) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    band = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(2.6))
    band.fill.solid(); band.fill.fore_color.rgb = TEAL; band.line.fill.background()
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.7), Inches(12), Inches(1.2))
    _text(tb.text_frame, d["title"], 32, True, RGBColor(0xFF, 0xFF, 0xFF))
    if d.get("subtitle"):
        tb2 = s.shapes.add_textbox(Inches(0.6), Inches(3.0), Inches(12), Inches(1))
        _text(tb2.text_frame, d["subtitle"], 18, color=GREY)
    _footer(s, prs, d, 1)
    for n, sl in enumerate(d.get("slides", []), 2):
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(12.3), Inches(0.8))
        _text(tb.text_frame, sl["title"], 24, True, TEAL)
        y = Inches(1.3)
        if sl.get("bullets"):
            body = s.shapes.add_textbox(Inches(0.6), y, Inches(12.1), Inches(4.5))
            tf = body.text_frame; tf.word_wrap = True
            for i, b in enumerate(sl["bullets"]):
                p = _text(tf, "• " + b, 14)
            y = y + Inches(0.4 * (len(sl["bullets"]) + 1))
        if sl.get("table"):
            cols, rows = sl["table"]["columns"], sl["table"].get("rows", [])
            shp = s.shapes.add_table(len(rows) + 1, len(cols), Inches(0.6), y, Inches(12.1), Inches(0.4) * (len(rows) + 1))
            t = shp.table
            for i, c in enumerate(cols):
                cell = t.cell(0, i); cell.text = str(c)
                cell.fill.solid(); cell.fill.fore_color.rgb = TEAL
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(11), True, RGBColor(0xFF, 0xFF, 0xFF)
            for ri, row in enumerate(rows, 1):
                for ci, v in enumerate(row[:len(cols)]):
                    cell = t.cell(ri, ci); cell.text = "" if v is None else str(v)
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.name, r.font.size = FONT, Pt(10)
        if sl.get("notes"):
            s.notes_slide.notes_text_frame.text = sl["notes"]
        _footer(s, prs, d, n)
    if d.get("sources"):
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(12.3), Inches(0.8))
        _text(tb.text_frame, "Sources", 24, True, TEAL)
        body = s.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(12.1), Inches(5))
        body.text_frame.word_wrap = True
        for i, src in enumerate(d["sources"], 1):
            _text(body.text_frame, f"{i}. {src}", 11)
        _footer(s, prs, d, len(prs.slides))
    prs.save(out)


if __name__ == "__main__":
    with open(sys.argv[1], encoding="utf-8") as fh:
        render(json.load(fh), sys.argv[2])
    print("OK", sys.argv[2])
