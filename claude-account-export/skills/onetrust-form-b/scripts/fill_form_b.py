# -*- coding: utf-8 -*-
"""
fill_form_b.py — write verified responses into a OneTrust Form B export xlsx.

The OneTrust 'Assessment Responses - v2' sheet is a round-trip import template:
only the Response and Justification columns may be edited; every other cell is
keyed to the question and must stay untouched, or OneTrust rejects the import.

Usage:
    python fill_form_b.py --xlsx <onetrust_export.xlsx> \\
                          --answers <answers.json> \\
                          --out <filled.xlsx>

answers.json: { "<question_id>": {"response": "...", "justification": "..."} }
e.g.          { "3.1": {"response": "1 - ISMS in place and documented with
                         continuous monitoring", "justification": "ISO 27001
                         cert attached"} }

The script edits the Response (col E) and Justification (col F) cells in place,
matching rows by the question-id prefix in column A. It validates single-select
answers against the 'Assessment Response Options' sheet and reports any answer
that is not an allowed option. It never modifies Question, Description,
Unique Identifier, Pre-selected Options or Question Type cells.
"""
import argparse, json, sys

def load_options(wb):
    """qid -> set of allowed option strings, from the options sheet."""
    opts = {}
    if 'Assessment Response Options' not in wb.sheetnames:
        return opts
    ws = wb['Assessment Response Options']
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return opts
    for ci, h in enumerate(rows[0]):
        if not h:
            continue
        qid = str(h).split()[0]
        vals = {str(r[ci]).strip() for r in rows[1:]
                if r[ci] is not None and str(r[ci]).strip()}
        opts[qid] = vals
    return opts

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', required=True, help='OneTrust Form B export xlsx')
    ap.add_argument('--answers', required=True, help='verified answers JSON')
    ap.add_argument('--out', required=True, help='output filled xlsx path')
    ap.add_argument('--sheet', default='Assessment Responses - v2')
    a = ap.parse_args()

    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl required: pip install openpyxl --break-system-packages")

    wb = openpyxl.load_workbook(a.xlsx)   # keep formulas/formatting
    if a.sheet not in wb.sheetnames:
        sys.exit(f"sheet '{a.sheet}' not found in {a.xlsx}")
    ws = wb[a.sheet]

    with open(a.answers, encoding='utf-8') as f:
        answers = json.load(f)

    options = load_options(wb)

    # locate the question-table header row (the row whose col A == 'Question')
    header_row = None
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 1).value).strip() == 'Question':
            header_row = r
            break
    if header_row is None:
        sys.exit("could not find the 'Question' header row")

    RESP_COL, JUST_COL = 5, 6   # E, F
    written, warnings, unmatched = 0, [], list(answers.keys())

    for r in range(header_row + 1, ws.max_row + 1):
        qcell = ws.cell(r, 1).value
        if not qcell:
            continue
        qid = str(qcell).split()[0]
        if qid not in answers:
            continue
        ans = answers[qid]
        resp = ans.get('response', '')
        just = ans.get('justification', '')
        # validate single-select answers against the option ladder
        if qid in options and options[qid] and resp:
            parts = [p.strip() for p in str(resp).split(',')]
            for p in parts:
                if p and p not in options[qid]:
                    warnings.append(
                        f"  {qid}: response {p!r} is not an allowed option")
        ws.cell(r, RESP_COL).value = resp
        ws.cell(r, JUST_COL).value = just
        written += 1
        if qid in unmatched:
            unmatched.remove(qid)

    wb.save(a.out)
    print(f"filled {written} responses -> {a.out}")
    if unmatched:
        print(f"WARNING - {len(unmatched)} answer id(s) not found in sheet: "
              + ', '.join(unmatched))
    if warnings:
        print("VALIDATION WARNINGS (answer not in option ladder):")
        print('\n'.join(warnings))
    if not unmatched and not warnings:
        print("OK - all answers matched and validated.")

if __name__ == '__main__':
    main()
